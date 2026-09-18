import logging
import os
import re
import socket
import time
from datetime import datetime
from urllib.parse import urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

from envio_email import enviar_relatorio
from login import login, URL as LOGIN_URL
from relatorio_execucao import (
    DESTINATARIO_PADRAO,
    Metricas,
    adquirir_lock_execucao,
    coletar_anexos,
    gerar_relatorio_html,
    liberar_lock_execucao,
    registrar_metricas,
    salvar_estatisticas,
    setup_logging,
    status_execucao,
)

_origin = urlparse(LOGIN_URL)
LOANS_URL = f"{_origin.scheme}://{_origin.netloc}/loans"

PASTA_PROJETO = os.path.dirname(os.path.abspath(__file__))
CAMINHO_RELATORIO = os.path.join(PASTA_PROJETO, "resultado_aprovacoes.xlsx")

VALOR_LIMITE = 10_000.00
MAX_PROPOSTAS = None  # None = extrai todas as propostas da lista
ITENS_POR_PAGINA = 50

logger = logging.getLogger("rpa_qualibank.propostas")


def _classificar_erro(e: Exception) -> str:
    """Classifica uma excecao para fins de log (nao influencia a logica de
    negocio, apenas o rotulo usado na mensagem registrada)."""
    if isinstance(e, PlaywrightTimeoutError):
        return "timeout"
    if isinstance(e, (AttributeError, ValueError, IndexError)):
        return "leitura"
    return "navegacao"


def parse_valor_brl(texto):
    texto = texto.strip().replace("R$", "").strip()
    texto = texto.replace(".", "").replace(",", ".")
    return float(texto)


def _cartao_da_secao(page, titulo):
    secao = page.locator(
        f'xpath=//div[contains(@class,"text-xl") and contains(@class,"mb-2") '
        f'and normalize-space(.)="{titulo}"]'
    )
    return secao.locator('xpath=following-sibling::div[contains(@class,"ajin-card")][1]')


def _campo(cartao, rotulo):
    campo = cartao.locator(
        f'xpath=.//*[contains(@class,"ajin-label")]'
        f'[normalize-space(text())="{rotulo}"]'
        f'/following-sibling::*[contains(@class,"ajin-value")][1]'
    )
    return campo.inner_text().strip()


def ler_proposta(page):
    cartao_proposta = _cartao_da_secao(page, "Proposta")
    contrato = _campo(cartao_proposta, "Contrato")
    valor = parse_valor_brl(_campo(cartao_proposta, "Valor do Contrato"))
    liquido = parse_valor_brl(_campo(cartao_proposta, "Líquido"))

    cartao_pessoais = _cartao_da_secao(page, "Dados Pessoais")
    nome = _campo(cartao_pessoais, "Name")

    cartao_log = _cartao_da_secao(page, "Log do Registro")
    data_proposta = _campo(cartao_log, "Data de Cadastro:")

    return contrato, nome, valor, liquido, data_proposta


def ler_data_aprovacao_promotora(linha):
    """Le a data/hora do status 'Aguardando Aprovação Promotora' direto na linha da lista."""
    nota = linha.locator('ajin-status-label[name="operationStatus"] .ajin-note').first
    if nota.count() == 0:
        return ""
    return nota.inner_text().strip()


def ler_loja(linha):
    """Le o nome da Loja direto na linha da lista (coluna Registro).
    Nao existe campo equivalente na pagina de detalhe da proposta."""
    valor = linha.locator(
        'xpath=.//span[contains(@class,"ajin-label")][normalize-space(text())="Loja:"]'
        '/following-sibling::span[contains(@class,"ajin-value")][1]'
    )
    if valor.count() == 0:
        return ""
    return valor.inner_text().strip()


def aprovar_proposta_real(page, contrato):
    """Se a aprovacao em si (Acoes -> Aprovacao Supervisor -> Confirmar) falhar,
    a excecao propaga para o chamador tratar. Uma falha so nos prints de
    evidencia (depois do Confirmar) NAO deve ser reportada como aprovacao
    falha, senao um contrato ja aprovado de verdade seria reprocessado."""
    logger.info(f"Iniciando aprovação real do contrato {contrato} (Ações -> Aprovação Supervisor).")
    page.click('button:has-text("Ações")')
    page.get_by_text("Aprovação Supervisor", exact=True).click()
    page.wait_for_timeout(500)
    page.fill("textarea[name='note']", "Aprovado via regra - RPA")
    page.get_by_role("button", name="Confirmar", exact=True).click()
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)
    logger.info(f"Aprovação confirmada no sistema para o contrato {contrato}.")

    try:
        page.screenshot(path=f"screenshots/aprovacao_{contrato}_proposta.png")
        page.get_by_text("Histórico", exact=True).click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(1000)
        page.screenshot(path=f"screenshots/aprovacao_{contrato}_historico.png")
    except Exception as e:
        logger.warning(f"Aprovação do contrato {contrato} concluída, mas falha ao capturar prints de evidência: {e}")


def _limpar_busca(page):
    """Garante que a lista nao fique presa a um filtro de busca residual
    (ex.: ficou preenchida com o nome do ultimo cliente visualizado)."""
    campo_busca = page.locator('input[placeholder="Pesquisar"]')
    if campo_busca.count() == 0:
        return
    if campo_busca.input_value():
        campo_busca.fill("")
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(500)


def _selecionar_todas_lojas(page):
    """Abre o filtro 'Loja', marca 'Selecionar Todos' e atualiza a lista,
    garantindo que a extracao enxergue propostas de todas as lojas."""
    botao_loja = page.locator(
        'xpath=//div[contains(@class,"text-base") and contains(@class,"text-secondary") '
        'and normalize-space(text())="Loja"]/following-sibling::button[@aria-haspopup="menu"][1]'
    )
    if botao_loja.count() == 0:
        return
    botao_loja.click()
    page.wait_for_timeout(500)

    page.get_by_role("menuitem", name="Selecionar Todos", exact=True).click()
    page.wait_for_timeout(500)

    _marcar_status_aguardando_supervisor(page)

    botao_refresh = page.locator('button:has(mat-icon[data-mat-icon-name="refresh"])').first
    botao_refresh.click()
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)


def _marcar_status_aguardando_supervisor(page):
    """Garante que o filtro de Status esteja marcado em 'Aguardando
    aprovação do Supervisor'. Esse filtro pode ficar desmarcado entre uma
    navegacao e outra, entao isso e checado (e corrigido, se preciso) toda
    vez que a lista de propostas e recarregada, nunca so uma vez no inicio.
    Se precisar corrigir, tambem atualiza a lista para refletir o filtro
    certo antes de continuar."""
    item = page.locator(
        'xpath=//span[contains(@class,"flex-auto") and contains(@class,"text-sm") '
        'and normalize-space(text())="Aguardando aprovação do Supervisor"]'
    )
    if item.count() == 0:
        return

    checkbox_input = item.locator('xpath=preceding-sibling::mat-checkbox[1]//input[@type="checkbox"]')
    if checkbox_input.count() == 0 or checkbox_input.is_checked():
        return

    item.click()
    page.wait_for_timeout(500)

    botao_refresh = page.locator('button:has(mat-icon[data-mat-icon-name="refresh"])').first
    if botao_refresh.count() > 0:
        botao_refresh.click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(1000)


def _total_propostas(page):
    page.locator("table.app-table-search tbody tr.cursor-pointer").first.wait_for(
        state="visible", timeout=30000
    )
    texto = page.locator("ajin-search-count").inner_text()
    numeros = re.findall(r"\d+", texto)
    total = int(numeros[-1]) if numeros else 0

    # Salvaguarda: a contagem do cabecalho pode nao ter carregado ainda;
    # nunca reportar menos do que ja esta renderizado na pagina atual.
    linhas_na_pagina = page.locator("table.app-table-search tbody tr.cursor-pointer").count()
    return max(total, linhas_na_pagina)


def _ir_proxima_pagina(page):
    botao = page.locator(
        'ajin-search-pagination button:has(mat-icon[data-mat-icon-name="arrow_forward_ios"])'
    )
    botao.click()
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(800)


MAX_TENTATIVAS = 5
# Protecao contra travamento: se nenhuma proposta for lida com sucesso por
# este tempo, ou se este numero de propostas seguidas falhar de vez, o
# processamento e interrompido (o relatorio e o e-mail ainda sao gerados).
TEMPO_MAX_SEM_PROGRESSO_SEGUNDOS = 300
MAX_FALHAS_CONSECUTIVAS = 2
ESPERA_ENTRE_TENTATIVAS_SEGUNDOS = 120


def _abrir_e_ler_proposta(page, i):
    pagina = i // ITENS_POR_PAGINA
    posicao = i % ITENS_POR_PAGINA

    page.goto(LOANS_URL)
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(800)
    _limpar_busca(page)
    _marcar_status_aguardando_supervisor(page)

    for _ in range(pagina):
        _ir_proxima_pagina(page)

    linha = page.locator("table.app-table-search tbody tr.cursor-pointer").nth(posicao)
    data_aprovacao_promotora = ler_data_aprovacao_promotora(linha)
    loja = ler_loja(linha)

    linha.locator('button[aria-haspopup="menu"]').first.click()
    page.wait_for_timeout(400)
    page.get_by_role("menuitem", name="Visualizar", exact=True).click()
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)

    contrato, nome, valor, liquido, data_proposta = ler_proposta(page)
    return contrato, nome, valor, liquido, data_proposta, data_aprovacao_promotora, loja


CHECKPOINT_A_CADA = 20


def processar_propostas(page, caminho_relatorio=None, sempre_anexar=False, metricas=None):
    """Se caminho_relatorio for informado, salva o relatorio periodicamente
    durante a execucao (a cada CHECKPOINT_A_CADA propostas), para nao perder
    o progresso caso o script seja interrompido no meio de um lote grande.
    Cada salvamento grava so as propostas novas desde o ultimo salvamento
    (nunca a lista inteira de novo), para nao duplicar linhas quando
    sempre_anexar=True.

    `metricas`, se informado (instancia de relatorio_execucao.Metricas), e
    apenas alimentado com dados que nao dá para derivar depois da lista de
    resultados (total encontrado, tentativas de reprocessamento) - nao afeta
    nenhuma decisao de aprovacao."""
    resultados = []
    ultimo_salvo = 0

    def salvar_novos():
        nonlocal ultimo_salvo
        if caminho_relatorio and len(resultados) > ultimo_salvo:
            gerar_relatorio(resultados[ultimo_salvo:], caminho_relatorio, sempre_anexar=sempre_anexar)
            ultimo_salvo = len(resultados)

    logger.info(f"Entrando na tela de propostas: {LOANS_URL}")
    page.goto(LOANS_URL)
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1000)
    _limpar_busca(page)
    _selecionar_todas_lojas(page)

    total = _total_propostas(page)
    n = total if MAX_PROPOSTAS is None else min(total, MAX_PROPOSTAS)
    if metricas is not None:
        metricas.total_encontradas = total
    logger.info(f"Propostas encontradas: {total}. Processando {n}.")

    ultimo_progresso = time.monotonic()
    falhas_consecutivas = 0

    def sem_progresso():
        return time.monotonic() - ultimo_progresso > TEMPO_MAX_SEM_PROGRESSO_SEGUNDOS

    try:
        for i in range(n):
            _processar_uma_proposta(page, i, resultados, metricas=metricas, deve_abortar=sem_progresso)
            if resultados[-1]["contrato"] == "?":
                falhas_consecutivas += 1
            else:
                falhas_consecutivas = 0
                ultimo_progresso = time.monotonic()

            motivo = None
            if falhas_consecutivas >= MAX_FALHAS_CONSECUTIVAS:
                motivo = f"{falhas_consecutivas} propostas seguidas falharam"
            elif sem_progresso():
                motivo = f"nenhuma proposta lida com sucesso há mais de {TEMPO_MAX_SEM_PROGRESSO_SEGUNDOS}s"
            if motivo and i + 1 < n:
                msg = f"Execução interrompida por proteção contra travamento: {motivo} (processadas {i + 1} de {n})."
                logger.error(msg)
                if metricas is not None:
                    metricas.erro_critico = msg
                break

            if (i + 1) % CHECKPOINT_A_CADA == 0:
                salvar_novos()
                logger.info(f"[CHECKPOINT] relatório salvo com {len(resultados)} propostas processadas até agora.")
    finally:
        salvar_novos()
        if caminho_relatorio:
            logger.info(f"[CHECKPOINT FINAL] relatório salvo com {len(resultados)} propostas.")

    return resultados


def _processar_uma_proposta(page, i, resultados, metricas=None, deve_abortar=None):
    erro = None
    dados = None
    for tentativa in range(1, MAX_TENTATIVAS + 1):
        try:
            dados = _abrir_e_ler_proposta(page, i)
            erro = None
            break
        except Exception as e:
            erro = e
            tipo_erro = _classificar_erro(e)
            logger.warning(f"[{i}] Tentativa {tentativa}/{MAX_TENTATIVAS} falhou ({tipo_erro}): {e}")
            if deve_abortar is not None and deve_abortar():
                logger.error(f"[{i}] Sem progresso há tempo demais; abandonando novas tentativas.")
                break
            if tentativa < MAX_TENTATIVAS:
                if metricas is not None:
                    metricas.tentativas_reprocessamento += 1
                logger.info(
                    f"    Aguardando {ESPERA_ENTRE_TENTATIVAS_SEGUNDOS}s antes de "
                    "tentar novamente (mesma sessao, sem novo login)..."
                )
                page.wait_for_timeout(ESPERA_ENTRE_TENTATIVAS_SEGUNDOS * 1000)

    if erro is not None:
        tipo_erro = _classificar_erro(erro)
        logger.error(f"[{i}] Desistindo após {MAX_TENTATIVAS} tentativas ({tipo_erro}): {erro}")
        resultados.append(
            {
                "contrato": "?",
                "nome": "",
                "loja": "",
                "valor": None,
                "liquido": None,
                "data_proposta": "",
                "data_aprovacao_promotora": "",
                "data_aprovacao_supervisor": "",
                "aprovado": False,
            }
        )
        return

    contrato, nome, valor, liquido, data_proposta, data_aprovacao_promotora, loja = dados
    logger.info(f"[{i}] Contrato {contrato} | Cliente: {nome} | Loja: {loja} | Valor Líquido: R$ {liquido:,.2f}")

    elegivel = liquido <= VALOR_LIMITE
    decisao = "APROVA" if elegivel else "PULA (valor > limite)"
    logger.info(f"[{i}] Contrato {contrato} -> {decisao}")

    data_aprovacao_supervisor = ""
    aprovado = False
    if elegivel:
        try:
            aprovar_proposta_real(page, contrato)
            data_aprovacao_supervisor = datetime.now().strftime("%d/%m/%Y %H:%M")
            aprovado = True
            logger.info(f"Aprovação realizada com sucesso para o contrato {contrato}.")
        except Exception as e:
            data_aprovacao_supervisor = f"ERRO: {e}"
            logger.error(f"Erro de aprovação no contrato {contrato}: {e}")
    else:
        logger.info(f"Proposta ignorada: contrato {contrato} (valor líquido acima do limite de R$ {VALOR_LIMITE:,.2f}).")

    resultados.append(
        {
            "contrato": contrato,
            "nome": nome,
            "loja": loja,
            "valor": valor,
            "liquido": liquido,
            "data_proposta": data_proposta,
            "data_aprovacao_promotora": data_aprovacao_promotora,
            "data_aprovacao_supervisor": data_aprovacao_supervisor,
            "aprovado": aprovado,
        }
    )


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
APROVADO_FILL = PatternFill("solid", fgColor="C6EFCE")
APROVADO_FONT = Font(color="006100")
SEM_FILL = PatternFill(fill_type=None)
FONTE_PADRAO = Font(color="000000")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
MOEDA_FORMATO = '"R$" #,##0.00'

COLUNAS_RELATORIO = [
    ("Código do Contrato", 22),
    ("Nome da Pessoa", 34),
    ("Valor Líquido", 18),
    ("Data e Horário da Proposta", 24),
    ("Data e Horário Aguardando Aprovação Promotora", 32),
    ("Data e Horário da Aprovação do Supervisor", 30),
    ("Decisão", 22),
    ("Loja", 26),
]


def _ano_da_proposta(item):
    """Ano usado para escolher a aba: ano da 'Data de Cadastro' real da proposta."""
    try:
        return datetime.strptime(item.get("data_proposta", ""), "%d/%m/%Y %H:%M").year
    except (TypeError, ValueError):
        return datetime.now().year


def _escrever_cabecalho(ws):
    """Escreve os cabecalhos de COLUNAS_RELATORIO que ainda nao existem na
    aba. Cobre tanto uma aba nova quanto uma aba criada por uma execucao
    anterior do script, que pode nao ter as colunas mais recentes."""
    for col_idx, (titulo, largura) in enumerate(COLUNAS_RELATORIO, start=1):
        if ws.cell(row=1, column=col_idx).value:
            continue
        celula = ws.cell(row=1, column=col_idx, value=titulo)
        celula.fill = HEADER_FILL
        celula.font = HEADER_FONT
        celula.alignment = Alignment(horizontal="center", vertical="center")
        celula.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = largura


def _obter_ou_criar_aba(wb, nome_aba):
    if nome_aba in wb.sheetnames:
        ws = wb[nome_aba]
        _escrever_cabecalho(ws)
        return ws

    # Reaproveita a aba padrao "Sheet" vazia da primeira criacao do workbook.
    if wb.sheetnames == ["Sheet"] and wb["Sheet"].max_row == 1 and wb["Sheet"]["A1"].value is None:
        ws = wb["Sheet"]
        ws.title = nome_aba
    else:
        ws = wb.create_sheet(nome_aba)

    ws.freeze_panes = "A2"
    _escrever_cabecalho(ws)
    return ws


def _escrever_linha(ws, row_idx, item):
    aprovado = item.get("aprovado", False)
    decisao_texto = "Aprovado" if aprovado else "Não Aprovado"

    valores = [
        item.get("contrato") or "-",
        item.get("nome") or "-",
        item.get("liquido"),
        item.get("data_proposta") or "-",
        item.get("data_aprovacao_promotora") or "",
        item.get("data_aprovacao_supervisor") or "",
        decisao_texto,
        item.get("loja") or "-",
    ]
    for col_idx, valor in enumerate(valores, start=1):
        celula = ws.cell(row=row_idx, column=col_idx, value=valor)
        celula.border = THIN_BORDER
        if col_idx == 3 and isinstance(valor, (int, float)):
            celula.number_format = MOEDA_FORMATO
        if col_idx in (4, 5, 6, 7):
            celula.alignment = Alignment(horizontal="center")
        celula.fill = APROVADO_FILL if aprovado else SEM_FILL
        celula.font = APROVADO_FONT if aprovado else FONTE_PADRAO


def gerar_relatorio(resultados, caminho, sempre_anexar=False):
    """Se sempre_anexar for True, cada proposta vira sempre uma linha nova
    no final da aba (nunca atualiza uma linha existente pelo contrato) -
    usado no relatorio real, para manter o historico de todas as execucoes."""

    if os.path.exists(caminho):
        wb = load_workbook(caminho)
    else:
        wb = Workbook()

    for item in resultados:
        nome_aba = str(_ano_da_proposta(item))
        ws = _obter_ou_criar_aba(wb, nome_aba)

        contrato = item.get("contrato") or "-"
        row_idx = None
        if not sempre_anexar and contrato != "-":
            for r in range(2, ws.max_row + 1):
                if ws.cell(row=r, column=1).value == contrato:
                    row_idx = r
                    break
        if row_idx is None:
            row_idx = ws.max_row + 1

        _escrever_linha(ws, row_idx, item)

    wb._sheets.sort(key=lambda ws: ws.title)
    wb.save(caminho)


def _executar_rpa(metricas: Metricas, caminho_relatorio: str, sempre_anexar: bool) -> list:
    """Executa o fluxo principal do RPA (login + processamento das
    propostas). Isolado em funcao propria para que o bloco __main__ possa
    envolve-lo num try/except/finally unico que sempre gera o relatorio e
    envia o e-mail, sem alterar a logica de aprovacao em si."""
    with sync_playwright() as p:
        browser, page = login(p)
        try:
            return processar_propostas(
                page,
                caminho_relatorio=caminho_relatorio,
                sempre_anexar=sempre_anexar,
                metricas=metricas,
            )
        finally:
            browser.close()


if __name__ == "__main__":
    os.makedirs(os.path.join(PASTA_PROJETO, "screenshots"), exist_ok=True)
    caminho_relatorio = CAMINHO_RELATORIO
    sempre_anexar = True

    _, caminho_log, timestamp_execucao = setup_logging(PASTA_PROJETO)

    # O agendador dispara uma nova execucao a cada 10 minutos, independente
    # da anterior ja ter terminado. Sem esse lock, uma execucao anterior
    # travada (ex.: aguardando MFA, pagina lenta) e a nova disparada em cima
    # dela tentariam logar ao mesmo tempo na mesma conta. Se ja houver uma
    # execucao rodando, esta encerra IMEDIATAMENTE, antes de qualquer
    # tentativa de login.
    if not adquirir_lock_execucao(PASTA_PROJETO):
        logger.warning(
            "Já existe uma execução do RPA em andamento neste momento - "
            "pulando este ciclo (nenhum login será tentado) para não rodar duas sessões em paralelo."
        )
        logging.shutdown()
        raise SystemExit(0)

    metricas = Metricas()
    resultados: list = []

    try:
        resultados = _executar_rpa(metricas, caminho_relatorio, sempre_anexar)
        logger.info(f"Relatório salvo em {caminho_relatorio} ({len(resultados)} propostas).")
    except Exception as e:
        metricas.erro_critico = str(e)
        logger.critical(f"Falha inesperada na execução do RPA: {e}", exc_info=True)
    finally:
        metricas.fim = datetime.now()
        stats = metricas.calcular(resultados)
        registrar_metricas(logger, stats)
        salvar_estatisticas(stats, PASTA_PROJETO, timestamp_execucao)

        anexos = coletar_anexos(caminho_log, caminho_relatorio, PASTA_PROJETO)
        corpo_html = gerar_relatorio_html(stats, resultados)
        _, status_texto_geral, _ = status_execucao(stats)
        assunto = (
            f"[RPA Qualibank] {status_texto_geral} — "
            f"{stats['fim_dt'].strftime('%d/%m/%Y %H:%M')} ({socket.gethostname()})"
        )

        try:
            metodo = enviar_relatorio(
                corpo_html, anexos, assunto=assunto, destinatario=DESTINATARIO_PADRAO,
                logger=logger,
            )
            logger.info(f"E-mail de relatório de execução enviado com sucesso (método: {metodo}).")
        except Exception as e:
            logger.error(f"Falha ao enviar e-mail de relatório (Outlook desktop e Web): {e}", exc_info=True)

        liberar_lock_execucao(PASTA_PROJETO)
        logging.shutdown()
